from datetime import datetime, time, date, timedelta
from decimal import Decimal
import csv
import io
from django.db import transaction
from django.utils import timezone
from django.contrib.auth.models import User
from django.db.models import Count, Q, Max, Avg, Sum
from crm.models import Lead, LeadFollowup, Customer, CustomerContact
from tenants.models import AuditEvent, CompanyMembership

@transaction.atomic
def convert_lead_to_customer(lead, user):
    """
    Converts a WON lead into a formal Customer record without duplicates.
    Preserves lead history.
    """
    company = lead.company

    # Check for duplicate customer by email or exact name match
    existing_customer = Customer.objects.filter(
        company=company,
        email__iexact=lead.email
    ).first()

    if not existing_customer and lead.company_name:
        existing_customer = Customer.objects.filter(
            company=company,
            name__iexact=lead.company_name
        ).first()

    if existing_customer:
        customer = existing_customer
        if not customer.lead_reference:
            customer.lead_reference = lead
            customer.save(update_fields=['lead_reference'])
    else:
        count = Customer.objects.filter(company=company).count() + 1
        cust_code = f"CUST-{1000 + count}"
        
        customer = Customer.objects.create(
            company=company,
            customer_code=cust_code,
            name=lead.company_name or lead.contact_person,
            company_name=lead.company_name,
            email=lead.email or f"info@{lead.company_name.lower().replace(' ', '')}.com",
            phone=lead.phone,
            billing_address=lead.address or "Address Pending",
            city="Kochi",
            state="Kerala",
            country="India",
            gstin="",
            lead_reference=lead
        )

        # Create primary contact
        CustomerContact.objects.create(
            customer=customer,
            name=lead.contact_person,
            email=lead.email,
            phone=lead.phone,
            designation="Primary Contact",
            is_primary=True
        )

    lead.status = 'WON'
    lead.pipeline_stage = 'Won'
    lead.conversion_probability = 100
    lead.save()

    AuditEvent.log(
        company=company,
        actor=user,
        action='CONVERT',
        module='crm',
        model_name='Lead',
        object_id=lead.id,
        object_repr=str(lead),
        changes={'converted_customer_id': customer.id, 'customer_code': customer.customer_code}
    )

    return customer


@transaction.atomic
def record_lead_followup(lead, staff, method, summary, outcome='', next_action='', next_followup_date=None, new_status=None):
    """
    Records a follow-up activity for a lead and syncs lead dates and status.
    """
    now = timezone.now()
    followup = LeadFollowup.objects.create(
        lead=lead,
        staff=staff,
        method=method,
        summary=summary,
        outcome=outcome,
        next_action=next_action,
        next_followup_date=next_followup_date
    )

    lead.last_followup_date = now
    
    if new_status and new_status in dict(Lead.STATUS_CHOICES):
        lead.status = new_status
        lead.pipeline_stage = dict(Lead.STATUS_CHOICES).get(new_status, new_status)
        if new_status == 'WON':
            lead.conversion_probability = 100
        elif new_status == 'LOST':
            lead.conversion_probability = 0
            next_followup_date = None
        elif new_status == 'QUALIFIED':
            lead.conversion_probability = 60
        elif new_status == 'PROPOSAL_SENT':
            lead.conversion_probability = 80

    if lead.status == 'LOST':
        lead.next_followup_date = None
        followup.next_followup_date = None
        followup.save(update_fields=['next_followup_date'])
    elif next_followup_date:
        lead.next_followup_date = next_followup_date

    lead.save()

    AuditEvent.log(
        company=lead.company,
        actor=staff,
        action='UPDATE',
        module='crm',
        model_name='LeadFollowup',
        object_id=str(lead.id),
        object_repr=f"{lead.lead_number} - Follow-up via {followup.get_method_display()} ({lead.company_name})",
        changes={
            'method': method,
            'summary': summary[:100],
            'outcome': outcome,
            'next_action': next_action,
            'next_followup_date': str(next_followup_date) if next_followup_date else None,
            'status': lead.status
        }
    )

    return followup


def assign_lead_to_staff(lead, assigned_user, actor):
    """
    Reassigns a single lead to a staff member.
    """
    old_assignee = lead.assigned_to.username if lead.assigned_to else "Unassigned"
    lead.assigned_to = assigned_user
    lead.save(update_fields=['assigned_to', 'updated_at'])

    AuditEvent.log(
        company=lead.company,
        actor=actor,
        action='UPDATE',
        module='crm',
        model_name='Lead',
        object_id=lead.id,
        object_repr=str(lead),
        changes={'reassigned_from': old_assignee, 'reassigned_to': assigned_user.username if assigned_user else "Unassigned"}
    )
    return lead


def bulk_assign_leads(company, lead_ids, assigned_user, actor):
    """
    Bulk assigns multiple leads to a staff member.
    """
    leads = Lead.objects.filter(company=company, id__in=lead_ids)
    count = leads.update(assigned_to=assigned_user)

    AuditEvent.log(
        company=company,
        actor=actor,
        action='UPDATE',
        module='crm',
        model_name='Lead',
        object_id='bulk',
        object_repr=f"Bulk assign {count} leads",
        changes={'lead_ids': list(lead_ids), 'assigned_to': assigned_user.username if assigned_user else "Unassigned"}
    )
    return count


def bulk_update_lead_status(company, lead_ids, new_status, actor):
    """
    Bulk updates the status of multiple leads.
    """
    valid_statuses = dict(Lead.STATUS_CHOICES)
    if new_status not in valid_statuses:
        return 0

    stage = valid_statuses[new_status]
    leads = Lead.objects.filter(company=company, id__in=lead_ids)
    if new_status == 'LOST':
        count = leads.update(status=new_status, pipeline_stage=stage, next_followup_date=None, conversion_probability=0)
    else:
        count = leads.update(status=new_status, pipeline_stage=stage)

    AuditEvent.log(
        company=company,
        actor=actor,
        action='UPDATE',
        module='crm',
        model_name='Lead',
        object_id='bulk',
        object_repr=f"Bulk update {count} leads status to {new_status}",
        changes={'lead_ids': list(lead_ids), 'status': new_status}
    )
    return count


def bulk_update_lead_priority(company, lead_ids, new_priority, actor):
    """
    Bulk updates the priority of multiple leads.
    """
    valid_priorities = dict(Lead.PRIORITY_CHOICES)
    if new_priority not in valid_priorities:
        return 0

    leads = Lead.objects.filter(company=company, id__in=lead_ids)
    count = leads.update(priority=new_priority)

    AuditEvent.log(
        company=company,
        actor=actor,
        action='UPDATE',
        module='crm',
        model_name='Lead',
        object_id='bulk',
        object_repr=f"Bulk update {count} leads priority to {new_priority}",
        changes={'lead_ids': list(lead_ids), 'priority': new_priority}
    )
    return count


def bulk_delete_leads(company, lead_ids, actor):
    """
    Bulk deletes multiple leads.
    """
    leads = Lead.objects.filter(company=company, id__in=lead_ids)
    count = leads.count()
    leads.delete()

    AuditEvent.log(
        company=company,
        actor=actor,
        action='DELETE',
        module='crm',
        model_name='Lead',
        object_id='bulk',
        object_repr=f"Bulk deleted {count} leads",
        changes={'lead_ids': list(lead_ids)}
    )
    return count


def generate_sample_leads_csv():
    """
    Generates a sample CSV template for bulk uploading leads.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Company Name",
        "Contact Person",
        "Phone",
        "Email",
        "Service Interested",
        "Estimated Value",
        "Priority",
        "Lead Source",
        "Address",
        "Notes",
        "Next Followup Date"
    ])
    writer.writerow([
        "Apex Digital Media",
        "Rahul Sharma",
        "+91 9876543210",
        "rahul@apexdigital.in",
        "Custom ERP & Cloud Solution",
        "250000.00",
        "HIGH",
        "Website",
        "MG Road, Kochi",
        "Looking for customized business automation",
        (timezone.now() + timedelta(days=2)).strftime("%Y-%m-%d 11:00")
    ])
    writer.writerow([
        "Green Valley Exports",
        "Ananya Nair",
        "+91 9845612345",
        "ananya@greenvalley.com",
        "Mobile App Development",
        "180000.00",
        "MEDIUM",
        "Referral",
        "Infopark Phase 2, Kochi",
        "Referred by policybazar client",
        (timezone.now() + timedelta(days=1)).strftime("%Y-%m-%d 15:30")
    ])
    writer.writerow([
        "Malabar Agro Logistics",
        "Mohammed Faizal",
        "+91 9745123456",
        "faizal@malabaragro.com",
        "Warehouse Management ERP",
        "350000.00",
        "URGENT",
        "Direct Outreach",
        "Calicut Beach Road, Kozhikode",
        "Urgent migration from legacy software",
        (timezone.now() + timedelta(days=3)).strftime("%Y-%m-%d 10:00")
    ])
    return output.getvalue()


def generate_sample_leads_excel():
    """
    Generates a professionally styled sample Excel (.xlsx) template for bulk uploading leads.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads Import Template"

    headers = [
        "Company Name",
        "Contact Person",
        "Phone",
        "Email",
        "Service Interested",
        "Estimated Value",
        "Priority",
        "Lead Source",
        "Address",
        "Notes",
        "Next Followup Date"
    ]
    ws.append(headers)

    # Style header row with distinctive dark fill and highlight mandatory fields
    header_fill = PatternFill(start_color="101720", end_color="101720", fill_type="solid")
    req_fill = PatternFill(start_color="008ec2", end_color="008ec2", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    for col_idx, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_name = headers[col_idx - 1]
        if col_name in ("Company Name", "Contact Person", "Phone"):
            cell.fill = req_fill
        else:
            cell.fill = header_fill

    sample_rows = [
        [
            "Apex Digital Media",
            "Rahul Sharma",
            "+91 9876543210",
            "rahul@apexdigital.in",
            "Custom ERP & Cloud Solution",
            "250000.00",
            "HIGH",
            "Website",
            "MG Road, Kochi",
            "Looking for customized business automation",
            (timezone.now() + timedelta(days=2)).strftime("%Y-%m-%d 11:00")
        ],
        [
            "Green Valley Exports",
            "Ananya Nair",
            "+91 9845612345",
            "ananya@greenvalley.com",
            "Mobile App Development",
            "180000.00",
            "MEDIUM",
            "Referral",
            "Infopark Phase 2, Kochi",
            "Referred by policybazar client",
            (timezone.now() + timedelta(days=1)).strftime("%Y-%m-%d 15:30")
        ],
        [
            "Malabar Agro Logistics",
            "Mohammed Faizal",
            "+91 9745123456",
            "faizal@malabaragro.com",
            "Warehouse Management ERP",
            "350000.00",
            "URGENT",
            "Direct Outreach",
            "Calicut Beach Road, Kozhikode",
            "Urgent migration from legacy software",
            (timezone.now() + timedelta(days=3)).strftime("%Y-%m-%d 10:00")
        ]
    ]

    for row_data in sample_rows:
        ws.append(row_data)

    for row in ws.iter_rows(min_row=2, max_row=len(sample_rows) + 1):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_and_import_leads(company, file_obj, uploaded_by, assignment_mode='self', target_user_id=None):
    """
    Parses and imports leads from CSV or XLSX Excel file.
    Phone number is strictly mandatory for each lead.
    """
    filename = file_obj.name.lower()
    rows = []

    if filename.endswith('.csv'):
        content = file_obj.read().decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(content))
        header = None
        for r in reader:
            if not r or not any(field.strip() for field in r):
                continue
            if header is None:
                header = [h.strip().lower() for h in r]
            else:
                row_dict = {}
                for idx, col_name in enumerate(header):
                    row_dict[col_name] = r[idx].strip() if idx < len(r) else ''
                rows.append(row_dict)
    elif filename.endswith('.xlsx') or filename.endswith('.xls'):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            sheet = wb.active
            header = None
            for r in sheet.iter_rows(values_only=True):
                if not r or not any(r):
                    continue
                if header is None:
                    header = [str(h or '').strip().lower() for h in r]
                else:
                    row_dict = {}
                    for idx, col_name in enumerate(header):
                        val = r[idx] if idx < len(r) else ''
                        if isinstance(val, float) and val.is_integer():
                            val = int(val)
                        elif isinstance(val, (datetime, date)):
                            val = val.strftime('%Y-%m-%d %H:%M')
                        row_dict[col_name] = str(val or '').strip()
                    rows.append(row_dict)
        except Exception as e:
            return 0, [f"Error reading Excel spreadsheet: {str(e)}. Please make sure it is a valid .xlsx file."]
    else:
        return 0, ["Unsupported file format. Please upload a .xlsx or .csv spreadsheet file."]

    if not rows:
        return 0, ["The uploaded spreadsheet file has no rows to import."]

    if not company:
        from tenants.models import Company
        company = Company.objects.filter(is_active=True).first() or Company.objects.first()
        if not company:
            company, _ = Company.objects.get_or_create(
                code="BYTEBOOT",
                defaults={
                    "name": "Byteboot Techno Solutions Pvt Ltd",
                    "currency_code": "INR",
                    "currency_symbol": "₹",
                    "is_active": True,
                }
            )

    # Resolve assignees
    company_staff_users = list(User.objects.filter(
        memberships__company=company,
        memberships__is_active=True
    ).distinct().order_by('username'))

    specific_user = None
    if assignment_mode == 'specific' and target_user_id:
        specific_user = User.objects.filter(id=target_user_id, memberships__company=company).first()

    created_count = 0
    errors = []
    
    current_max_count = Lead.objects.filter(company=company).count()

    def get_val(row_dict, aliases, default=''):
        for alias in aliases:
            for k in row_dict:
                if alias in k:
                    v = row_dict[k]
                    if v:
                        return v
        return default

    round_robin_idx = 0

    with transaction.atomic():
        for line_num, row in enumerate(rows, start=2):
            company_name = get_val(row, ['company name', 'company', 'client name', 'client'])
            contact_person = get_val(row, ['contact person', 'contact', 'person', 'name'])
            phone = get_val(row, ['phone', 'mobile', 'cell', 'tel', 'contact number', 'phone number'])
            email = get_val(row, ['email', 'mail'])
            service = get_val(row, ['service interested', 'service', 'product', 'requirement'], 'Software Development')
            value_str = get_val(row, ['estimated value', 'estimated', 'value', 'amount', 'budget'], '0.00')
            priority = get_val(row, ['priority'], 'MEDIUM').upper()
            source = get_val(row, ['lead source', 'source'], 'Website')
            address = get_val(row, ['address', 'location'], '')
            notes = get_val(row, ['notes', 'note', 'remark', 'remarks'], '')
            fup_date_str = get_val(row, ['next followup date', 'followup date', 'followup', 'follow up', 'next follow'], '')

            if not company_name and not contact_person:
                errors.append(f"Row {line_num}: Missing both Company Name and Contact Person. Skipped.")
                continue

            # Strict mandatory Phone number validation
            if not phone or not phone.strip() or phone.strip().lower() in ('none', 'null', 'nan', 'not provided', '-', 'n/a'):
                lead_ref = company_name or contact_person or f"Row {line_num}"
                errors.append(f"Row {line_num} ({lead_ref}): Phone number is mandatory. Skipped.")
                continue

            phone = phone.strip()
            if phone.endswith('.0'):
                phone = phone[:-2]

            # Parse estimated value
            try:
                # clean currency symbols
                clean_val = value_str.replace('₹', '').replace('$', '').replace(',', '').strip()
                estimated_val = Decimal(clean_val) if clean_val else Decimal('0.00')
            except Exception:
                estimated_val = Decimal('0.00')

            # Normalize priority
            if priority not in ('LOW', 'MEDIUM', 'HIGH', 'URGENT'):
                priority = 'MEDIUM'

            # Parse next followup date
            next_fup = None
            if fup_date_str:
                for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%d-%m-%Y %H:%M', '%d-%m-%Y'):
                    try:
                        parsed_dt = datetime.strptime(fup_date_str, fmt)
                        if timezone.is_naive(parsed_dt):
                            next_fup = timezone.make_aware(parsed_dt)
                        else:
                            next_fup = parsed_dt
                        break
                    except (ValueError, TypeError):
                        pass

            # Determine assignment for this row
            row_assigned_to = None
            if assignment_mode == 'self':
                row_assigned_to = uploaded_by
            elif assignment_mode == 'specific':
                row_assigned_to = specific_user or uploaded_by
            elif assignment_mode == 'round_robin':
                if company_staff_users:
                    row_assigned_to = company_staff_users[round_robin_idx % len(company_staff_users)]
                    round_robin_idx += 1
                else:
                    row_assigned_to = uploaded_by
            elif assignment_mode == 'unassigned':
                row_assigned_to = None
            else:
                row_assigned_to = uploaded_by

            current_max_count += 1
            lead_num = f"LD-{10000 + current_max_count}"

            # Check if lead number already taken, find next safe
            while Lead.objects.filter(company=company, lead_number=lead_num).exists():
                current_max_count += 1
                lead_num = f"LD-{10000 + current_max_count}"

            lead = Lead.objects.create(
                company=company,
                lead_number=lead_num,
                company_name=company_name or contact_person,
                contact_person=contact_person or company_name,
                phone=phone,
                email=email,
                address=address,
                source=source,
                service_interested=service,
                estimated_value=estimated_val,
                priority=priority,
                status='NEW',
                pipeline_stage='New',
                notes=notes,
                next_followup_date=next_fup,
                created_by=uploaded_by,
                assigned_to=row_assigned_to
            )
            created_count += 1

    AuditEvent.log(
        company=company,
        actor=uploaded_by,
        action='CREATE',
        module='crm',
        model_name='Lead',
        object_id='bulk_upload',
        object_repr=f"Imported {created_count} leads via file {file_obj.name}",
        changes={'imported_count': created_count, 'assignment_mode': assignment_mode}
    )

    return created_count, errors


def get_staff_crm_performance(company, start_date=None, end_date=None, sort_by='-followups_count', staff_id=None):
    """
    Aggregates performance analytics for each staff member in the company:
    - Followups logged in date range
    - Calls made (PHONE)
    - Converted leads / Won leads
    - Total assigned leads
    - Overdue followups pending
    """
    now = timezone.localtime(timezone.now())
    today = now.date()
    start_of_today = timezone.make_aware(datetime.combine(today, time.min))

    # Resolve date boundaries
    if not start_date:
        # Default to first day of current month
        start_date = today.replace(day=1)
    if not end_date:
        end_date = today

    range_start = timezone.make_aware(datetime.combine(start_date, time.min))
    range_end = timezone.make_aware(datetime.combine(end_date, time.max))

    # Get company members
    memberships = CompanyMembership.objects.filter(
        company=company,
        is_active=True
    ).select_related('user', 'user__employee_profile', 'custom_role')

    if staff_id:
        memberships = memberships.filter(user_id=staff_id)

    staff_stats = []
    
    total_company_followups = 0
    total_company_calls = 0
    total_company_converted = 0

    for m in memberships:
        user = m.user
        emp = getattr(user, 'employee_profile', None)
        full_name = emp.full_name if emp else (user.get_full_name() or user.username)
        designation = emp.designation if emp else (m.custom_role.name if m.custom_role else m.get_role_display())

        # Followups in period by this staff
        user_followups = LeadFollowup.objects.filter(
            lead__company=company,
            staff=user,
            interaction_time__range=(range_start, range_end)
        )
        followups_count = user_followups.count()
        calls_count = user_followups.filter(method='PHONE').count()
        whatsapp_count = user_followups.filter(method='WHATSAPP').count()
        meetings_count = user_followups.filter(method='MEETING').count()
        emails_count = user_followups.filter(method='EMAIL').count()

        # Assigned leads
        assigned_leads = Lead.objects.filter(company=company, assigned_to=user)
        total_assigned = assigned_leads.count()

        # Won/Converted leads in period
        converted_count = assigned_leads.filter(status='WON', updated_at__range=(range_start, range_end)).count()
        lost_count = assigned_leads.filter(status='LOST', updated_at__range=(range_start, range_end)).count()

        # Overdue followups pending
        overdue_count = assigned_leads.filter(
            next_followup_date__lt=start_of_today
        ).exclude(status__in=['WON', 'LOST']).count()

        # Today followups
        today_fup_count = assigned_leads.filter(
            next_followup_date__range=(start_of_today, timezone.make_aware(datetime.combine(today, time.max)))
        ).exclude(status__in=['WON', 'LOST']).count()

        # Last activity
        last_activity = user_followups.order_by('-interaction_time').first()
        last_activity_time = last_activity.interaction_time if last_activity else None

        conversion_rate = round((converted_count / total_assigned * 100), 1) if total_assigned > 0 else 0.0

        total_company_followups += followups_count
        total_company_calls += calls_count
        total_company_converted += converted_count

        staff_stats.append({
            'user': user,
            'full_name': full_name,
            'username': user.username,
            'email': user.email,
            'designation': designation,
            'role': m.role,
            'followups_count': followups_count,
            'calls_count': calls_count,
            'whatsapp_count': whatsapp_count,
            'meetings_count': meetings_count,
            'emails_count': emails_count,
            'total_assigned': total_assigned,
            'converted_count': converted_count,
            'lost_count': lost_count,
            'overdue_count': overdue_count,
            'today_fup_count': today_fup_count,
            'last_activity_time': last_activity_time,
            'conversion_rate': conversion_rate
        })

    # Sort stats
    reverse = not sort_by.startswith('+')
    sort_key = sort_by.lstrip('+-')
    if sort_key == 'name':
        staff_stats.sort(key=lambda s: s['full_name'].lower(), reverse=reverse)
    elif sort_key in ('followups_count', 'calls_count', 'converted_count', 'total_assigned', 'conversion_rate', 'overdue_count'):
        staff_stats.sort(key=lambda s: s.get(sort_key, 0), reverse=reverse)
    else:
        staff_stats.sort(key=lambda s: s['followups_count'], reverse=True)

    top_performer = max(staff_stats, key=lambda s: s['followups_count']) if staff_stats and staff_stats[0]['followups_count'] > 0 else None

    return {
        'staff_stats': staff_stats,
        'summary': {
            'total_followups': total_company_followups,
            'total_calls': total_company_calls,
            'total_converted': total_company_converted,
            'total_staff': len(staff_stats),
            'top_performer': top_performer,
            'start_date': start_date,
            'end_date': end_date
        }
    }


def get_crm_360_analytics_data(company, start_date=None, end_date=None, staff_id=None):
    """
    Computes 360-degree comprehensive CRM performance analytics across multiple business dimensions:
    1. Executive KPIs (Outreach, Deals Won, Win Rate, Pipeline Value, Won Revenue, Overdue tasks)
    2. Staff Outreach Velocity & Conversion Scorecard
    3. Lead Quality & Acquisition Channels Matrix (Volume, Win Rate %, Won Value, Quality Tier)
    4. Sales Pipeline & Conversion Funnel Progression (Stage breakdown, Drop-offs, Lost deal reasons)
    5. Follow-ups & Activity Audit Trail (Chronological interaction log)
    """
    now = timezone.localtime(timezone.now())
    today = now.date()
    start_of_today = timezone.make_aware(datetime.combine(today, time.min))
    end_of_today = timezone.make_aware(datetime.combine(today, time.max))

    # Resolve date boundaries
    if not start_date:
        start_date = today.replace(day=1)
    if not end_date:
        end_date = today

    range_start = timezone.make_aware(datetime.combine(start_date, time.min))
    range_end = timezone.make_aware(datetime.combine(end_date, time.max))

    # Staff user filter (if specified)
    selected_staff = None
    if staff_id:
        selected_staff = User.objects.filter(id=staff_id).first()

    # Base queries scoped by company
    all_leads = Lead.objects.filter(company=company)
    all_fups = LeadFollowup.objects.filter(lead__company=company, interaction_time__range=(range_start, range_end))

    # Apply staff filtering if selected
    scoped_leads = all_leads.filter(assigned_to=selected_staff) if selected_staff else all_leads
    scoped_fups = all_fups.filter(staff=selected_staff) if selected_staff else all_fups

    # Company Memberships
    memberships = CompanyMembership.objects.filter(
        company=company,
        is_active=True
    ).select_related('user', 'user__employee_profile', 'custom_role')

    if selected_staff:
        memberships = memberships.filter(user=selected_staff)

    # 1. Staff Performance Scorecards
    staff_stats = []
    total_company_followups = 0
    total_company_calls = 0
    total_company_converted = 0
    total_company_won_val = Decimal('0.00')

    for m in memberships:
        user = m.user
        emp = getattr(user, 'employee_profile', None)
        full_name = emp.full_name if emp else (user.get_full_name() or user.username)
        designation = emp.designation if emp else (m.custom_role.name if m.custom_role else m.get_role_display())

        # Followups in period by this staff
        user_fups = all_fups.filter(staff=user)
        fup_cnt = user_fups.count()
        calls_cnt = user_fups.filter(method='PHONE').count()
        wa_cnt = user_fups.filter(method='WHATSAPP').count()
        meet_cnt = user_fups.filter(method='MEETING').count()
        email_cnt = user_fups.filter(method='EMAIL').count()
        video_cnt = user_fups.filter(method='VIDEO').count()

        # Assigned leads
        user_leads = all_leads.filter(assigned_to=user)
        total_assigned = user_leads.count()

        # Converted & Lost leads in period
        won_leads = user_leads.filter(status='WON', updated_at__range=(range_start, range_end))
        won_cnt = won_leads.count()
        lost_leads = user_leads.filter(status='LOST', updated_at__range=(range_start, range_end))
        lost_cnt = lost_leads.count()

        # Revenue and pipeline values
        won_rev = sum((l.estimated_value or Decimal('0.00')) for l in won_leads)
        active_user_leads = user_leads.exclude(status__in=['WON', 'LOST'])
        pipeline_val = sum((l.estimated_value or Decimal('0.00')) for l in active_user_leads)

        # Overdue & today followups
        overdue_cnt = active_user_leads.filter(next_followup_date__lt=start_of_today).count()
        today_cnt = active_user_leads.filter(next_followup_date__range=(start_of_today, end_of_today)).count()

        # Conversion rate
        conv_rate = round((won_cnt / total_assigned * 100), 1) if total_assigned > 0 else 0.0

        # Last activity
        last_fup = user_fups.order_by('-interaction_time').first()
        last_time = last_fup.interaction_time if last_fup else None

        total_company_followups += fup_cnt
        total_company_calls += calls_cnt
        total_company_converted += won_cnt
        total_company_won_val += won_rev

        staff_stats.append({
            'user_id': user.id,
            'user': user,
            'full_name': full_name,
            'username': user.username,
            'email': user.email,
            'designation': designation,
            'role': m.role,
            'total_assigned': total_assigned,
            'followups_count': fup_cnt,
            'calls_count': calls_cnt,
            'whatsapp_count': wa_cnt,
            'meetings_count': meet_cnt,
            'emails_count': email_cnt,
            'videos_count': video_cnt,
            'converted_count': won_cnt,
            'lost_count': lost_cnt,
            'conversion_rate': conv_rate,
            'won_revenue': won_rev,
            'pipeline_value': pipeline_val,
            'overdue_count': overdue_cnt,
            'today_fup_count': today_cnt,
            'last_activity_time': last_time,
        })

    staff_stats.sort(key=lambda s: s['followups_count'], reverse=True)

    # 2. Lead Quality & Acquisition Sources Analysis
    raw_sources = set(scoped_leads.values_list('source', flat=True))
    sources_list = sorted([s for s in raw_sources if s])
    if not sources_list:
        sources_list = ['Website', 'Referral', 'Social Media', 'Cold Call']

    lead_sources = []
    total_leads_count = scoped_leads.count()

    for src in sources_list:
        src_leads = scoped_leads.filter(source=src)
        s_total = src_leads.count()
        s_won_qs = src_leads.filter(status='WON', updated_at__range=(range_start, range_end))
        s_won = s_won_qs.count()
        s_lost_qs = src_leads.filter(status='LOST', updated_at__range=(range_start, range_end))
        s_lost = s_lost_qs.count()
        s_open = src_leads.exclude(status__in=['WON', 'LOST']).count()

        s_val = sum((l.estimated_value or Decimal('0.00')) for l in src_leads)
        s_won_val = sum((l.estimated_value or Decimal('0.00')) for l in s_won_qs)
        s_rate = round((s_won / s_total * 100), 1) if s_total > 0 else 0.0
        s_avg_deal = round(s_val / s_total, 2) if s_total > 0 else Decimal('0.00')
        s_high_prio = src_leads.filter(priority__in=['URGENT', 'HIGH']).count()

        # Quality tier assignment
        if s_rate >= 25.0 or s_won_val >= Decimal('200000.00'):
            tier = 'Grade A (High Win Rate)'
        elif s_rate >= 10.0 or s_high_prio >= 2 or s_total >= 3:
            tier = 'Grade B (Good Potential)'
        else:
            tier = 'Grade C (Low Conversion)'

        lead_sources.append({
            'source': src,
            'total_leads': s_total,
            'won_count': s_won,
            'lost_count': s_lost,
            'open_count': s_open,
            'conversion_rate': s_rate,
            'total_value': s_val,
            'won_value': s_won_val,
            'avg_deal_size': s_avg_deal,
            'high_priority_count': s_high_prio,
            'quality_tier': tier
        })

    lead_sources.sort(key=lambda s: s['total_value'], reverse=True)

    # 3. Priority Distribution Matrix
    priority_matrix = []
    for p_code, p_name in Lead.PRIORITY_CHOICES:
        p_leads = scoped_leads.filter(priority=p_code)
        p_cnt = p_leads.count()
        p_won = p_leads.filter(status='WON', updated_at__range=(range_start, range_end)).count()
        p_rate = round((p_won / p_cnt * 100), 1) if p_cnt > 0 else 0.0
        p_val = sum((l.estimated_value or Decimal('0.00')) for l in p_leads)
        priority_matrix.append({
            'code': p_code,
            'label': p_name,
            'total_count': p_cnt,
            'won_count': p_won,
            'conversion_rate': p_rate,
            'total_value': p_val
        })

    # 4. Pipeline Stages & Conversion Funnel (Both Period-Filtered Activity & Cumulative All-Time Static)
    pipeline_funnel = []
    for st_code, st_name in Lead.STATUS_CHOICES:
        # All-time static cumulative leads in this stage
        all_st_leads = scoped_leads.filter(status=st_code)
        all_cnt = all_st_leads.count()
        all_val = sum((l.estimated_value or Decimal('0.00')) for l in all_st_leads)

        # Date-wise period filtered leads (leads created, updated, or converted in selected date window)
        if st_code == 'WON':
            period_st_leads = scoped_leads.filter(status='WON', updated_at__range=(range_start, range_end))
        elif st_code == 'LOST':
            period_st_leads = scoped_leads.filter(status='LOST', updated_at__range=(range_start, range_end))
        elif st_code == 'NEW':
            period_st_leads = scoped_leads.filter(status='NEW', created_at__range=(range_start, range_end))
        else:
            period_st_leads = scoped_leads.filter(
                status=st_code
            ).filter(
                Q(updated_at__range=(range_start, range_end)) |
                Q(created_at__range=(range_start, range_end)) |
                Q(followups__interaction_time__range=(range_start, range_end))
            ).distinct()

        period_cnt = period_st_leads.count()
        period_val = sum((l.estimated_value or Decimal('0.00')) for l in period_st_leads)

        st_pct = round((all_cnt / total_leads_count * 100), 1) if total_leads_count > 0 else 0.0
        avg_prob = round(all_st_leads.aggregate(avg=Avg('conversion_probability'))['avg'] or 0, 1) if all_cnt > 0 else 0.0

        pipeline_funnel.append({
            'status_code': st_code,
            'status_name': st_name,
            'count': all_cnt,
            'pipeline_value': all_val,
            'percentage': st_pct,
            'period_count': period_cnt,
            'period_value': period_val,
            'avg_probability': avg_prob
        })

    # 5. Lost Deals Analysis
    lost_reasons = []
    lost_qs = scoped_leads.filter(status='LOST')
    for l in lost_qs:
        reason_text = (l.lost_reason or '').strip() or 'No specific reason specified'
        existing = next((r for r in lost_reasons if r['reason'] == reason_text), None)
        val = l.estimated_value or Decimal('0.00')
        if existing:
            existing['count'] += 1
            existing['lost_value'] += val
        else:
            lost_reasons.append({
                'reason': reason_text,
                'count': 1,
                'lost_value': val
            })
    lost_reasons.sort(key=lambda r: r['count'], reverse=True)

    # 6. Granular Follow-ups Activity Trail (up to 200 entries)
    activity_logs = []
    for f in scoped_fups.select_related('staff', 'lead').order_by('-interaction_time')[:200]:
        staff_user = f.staff
        s_emp = getattr(staff_user, 'employee_profile', None)
        s_name = s_emp.full_name if s_emp else (staff_user.get_full_name() or staff_user.username)
        activity_logs.append({
            'timestamp': f.interaction_time,
            'staff_name': s_name,
            'lead_number': f.lead.lead_number,
            'company_name': f.lead.company_name,
            'contact_person': f.lead.contact_person,
            'method': f.get_method_display(),
            'summary': f.summary or '',
            'outcome': f.outcome or 'Pending next step',
            'next_action': f.next_action or 'N/A',
            'next_followup_date': f.next_followup_date
        })

    # 7. Overall Summary KPIs
    active_leads_qs = scoped_leads.exclude(status__in=['WON', 'LOST'])
    won_leads_period_qs = scoped_leads.filter(status='WON', updated_at__range=(range_start, range_end))
    lost_leads_period_qs = scoped_leads.filter(status='LOST', updated_at__range=(range_start, range_end))

    total_pipeline_val = sum((l.estimated_value or Decimal('0.00')) for l in active_leads_qs)
    total_won_revenue = sum((l.estimated_value or Decimal('0.00')) for l in won_leads_period_qs)
    total_won_cnt = won_leads_period_qs.count()
    total_lost_cnt = lost_leads_period_qs.count()

    overall_conv_rate = round((total_won_cnt / total_leads_count * 100), 1) if total_leads_count > 0 else 0.0
    closed_deals = total_won_cnt + total_lost_cnt
    pipeline_win_rate = round((total_won_cnt / closed_deals * 100), 1) if closed_deals > 0 else 0.0
    avg_deal_size = round(total_won_revenue / total_won_cnt, 2) if total_won_cnt > 0 else (round(total_pipeline_val / total_leads_count, 2) if total_leads_count > 0 else Decimal('0.00'))

    overdue_total = active_leads_qs.filter(next_followup_date__lt=start_of_today).count()
    today_fup_total = active_leads_qs.filter(next_followup_date__range=(start_of_today, end_of_today)).count()

    top_performer_fups = max(staff_stats, key=lambda s: s['followups_count']) if staff_stats and any(s['followups_count'] > 0 for s in staff_stats) else None
    top_performer_conv = max(staff_stats, key=lambda s: s['converted_count']) if staff_stats and any(s['converted_count'] > 0 for s in staff_stats) else None
    top_performer_rev = max(staff_stats, key=lambda s: s['won_revenue']) if staff_stats and any(s['won_revenue'] > 0 for s in staff_stats) else None

    return {
        'company': company,
        'period_info': {
            'start_date': start_date,
            'end_date': end_date,
            'range_start': range_start,
            'range_end': range_end,
            'selected_staff': selected_staff,
            'generated_at': now
        },
        'summary_kpis': {
            'total_leads': total_leads_count,
            'new_leads_period': scoped_leads.filter(created_at__range=(range_start, range_end)).count(),
            'total_followups': scoped_fups.count(),
            'total_calls': scoped_fups.filter(method='PHONE').count(),
            'total_whatsapp': scoped_fups.filter(method='WHATSAPP').count(),
            'total_meetings': scoped_fups.filter(method='MEETING').count(),
            'total_emails': scoped_fups.filter(method='EMAIL').count(),
            'total_videos': scoped_fups.filter(method='VIDEO').count(),
            'total_converted': total_won_cnt,
            'total_lost': total_lost_cnt,
            'overall_conversion_rate': overall_conv_rate,
            'pipeline_win_rate': pipeline_win_rate,
            'active_pipeline_value': total_pipeline_val,
            'won_revenue': total_won_revenue,
            'avg_deal_size': avg_deal_size,
            'overdue_count': overdue_total,
            'today_fup_count': today_fup_total,
            'top_staff_followups': top_performer_fups,
            'top_staff_converted': top_performer_conv,
            'top_staff_revenue': top_performer_rev,
            'total_staff': len(staff_stats)
        },
        'staff_performance': staff_stats,
        'lead_sources': lead_sources,
        'priority_matrix': priority_matrix,
        'pipeline_funnel': pipeline_funnel,
        'lost_reasons': lost_reasons,
        'activity_logs': activity_logs
    }


def generate_crm_excel_report(company, report_type='360_master', start_date=None, end_date=None, staff_id=None):
    """
    Generates an executive, beautifully formatted multi-sheet Excel (.xlsx) workbook for CRM Analytics.
    Supported report_type:
    - '360_master': Complete 360° master report (Summary, Staff, Sources, Funnel, Activities)
    - 'staff_performance': Staff performance scorecard & outreach velocity
    - 'lead_quality': Lead quality channels & priority matrix
    - 'conversion_funnel': Sales pipeline funnel & conversion velocity
    - 'activity_log': Detailed chronological follow-up audit trail
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    data = get_crm_360_analytics_data(company, start_date=start_date, end_date=end_date, staff_id=staff_id)
    kpis = data['summary_kpis']
    period = data['period_info']

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    # Styling Palettes
    header_dark_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_cyan_fill = PatternFill(start_color="00A0DA", end_color="00A0DA", fill_type="solid")
    sub_header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    zebra_even_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    zebra_odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    kpi_box_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    total_row_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    font_title = Font(name="Arial", size=14, bold=True, color="0F172A")
    font_subtitle = Font(name="Arial", size=9, italic=True, color="64748B")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_cell = Font(name="Arial", size=9, color="1E293B")
    font_cell_bold = Font(name="Arial", size=9, bold=True, color="0F172A")
    font_kpi_num = Font(name="Arial", size=14, bold=True, color="00A0DA")
    font_kpi_label = Font(name="Arial", size=8, bold=True, color="64748B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    thick_bottom = Border(bottom=Side(style='medium', color='00A0DA'))

    def auto_fit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    s_val = str(cell.value)
                    if len(s_val) > max_len:
                        max_len = len(s_val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 13), 50)

    # ------------------ SHEET: Executive Overview ------------------
    if report_type in ('360_master', 'staff_performance', 'lead_quality', 'conversion_funnel'):
        ws_summary = wb.create_sheet(title="Executive Overview")
        
        # Banner Title
        ws_summary["A1"] = f"{company.name.upper()} — CRM 360° INTELLIGENCE AUDIT"
        ws_summary["A1"].font = font_title
        ws_summary["A2"] = f"Period: {period['start_date'].strftime('%d %b %Y')} to {period['end_date'].strftime('%d %b %Y')} | Generated: {period['generated_at'].strftime('%d %b %Y %H:%M')} | Scope: {'All Staff' if not period['selected_staff'] else period['selected_staff'].get_full_name()}"
        ws_summary["A2"].font = font_subtitle

        # KPI Block Table
        kpi_metrics = [
            ("TOTAL OUTREACH", kpis['total_followups'], "Total Follow-ups Logged"),
            ("DIRECT PHONE CALLS", kpis['total_calls'], "Voice Conversations"),
            ("DEALS CONVERTED (WON)", kpis['total_converted'], "Paying Customers Won"),
            ("WON REVENUE (₹)", float(kpis['won_revenue']), "Total Won Deal Value"),
            ("ACTIVE PIPELINE (₹)", float(kpis['active_pipeline_value']), "Open Deals in Funnel"),
            ("OVERALL WIN RATE", f"{kpis['overall_conversion_rate']}%", "Leads Won / Total")
        ]

        row_idx = 4
        ws_summary.cell(row=row_idx, column=1, value="KEY EXECUTIVE PERFORMANCE INDICATORS").font = font_cell_bold
        row_idx += 1

        for idx, (label, val, desc) in enumerate(kpi_metrics, 1):
            c1 = ws_summary.cell(row=row_idx, column=idx, value=label)
            c1.font = font_kpi_label
            c1.fill = sub_header_fill
            c1.font = Font(name="Arial", size=8, bold=True, color="FFFFFF")
            c1.alignment = Alignment(horizontal="center")

            c2 = ws_summary.cell(row=row_idx + 1, column=idx, value=val)
            c2.font = font_kpi_num
            c2.fill = kpi_box_fill
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c2.border = thin_border
            if isinstance(val, float):
                c2.number_format = '₹ #,##0.00'

            c3 = ws_summary.cell(row=row_idx + 2, column=idx, value=desc)
            c3.font = Font(name="Arial", size=7, italic=True, color="64748B")
            c3.fill = kpi_box_fill
            c3.alignment = Alignment(horizontal="center")
            c3.border = thin_border

        # Top Performers Highlight
        row_idx += 5
        ws_summary.cell(row=row_idx, column=1, value="EXECUTIVE LEADERBOARD & RECOGNITION").font = font_cell_bold
        row_idx += 1
        ws_summary.cell(row=row_idx, column=1, value="Metric Category").fill = header_dark_fill
        ws_summary.cell(row=row_idx, column=1).font = font_header
        ws_summary.cell(row=row_idx, column=2, value="Top Performer").fill = header_dark_fill
        ws_summary.cell(row=row_idx, column=2).font = font_header
        ws_summary.cell(row=row_idx, column=3, value="Performance Score").fill = header_dark_fill
        ws_summary.cell(row=row_idx, column=3).font = font_header

        leaders = [
            ("Highest Outreach Velocity", kpis['top_staff_followups']['full_name'] if kpis['top_staff_followups'] else 'None', f"{kpis['top_staff_followups']['followups_count']} follow-ups" if kpis['top_staff_followups'] else '-'),
            ("Most Deals Converted (Won)", kpis['top_staff_converted']['full_name'] if kpis['top_staff_converted'] else 'None', f"{kpis['top_staff_converted']['converted_count']} won deals" if kpis['top_staff_converted'] else '-'),
            ("Highest Won Revenue Closed", kpis['top_staff_revenue']['full_name'] if kpis['top_staff_revenue'] else 'None', f"₹ {float(kpis['top_staff_revenue']['won_revenue']):,.2f}" if kpis['top_staff_revenue'] else '-')
        ]

        for cat, name, score in leaders:
            row_idx += 1
            ws_summary.cell(row=row_idx, column=1, value=cat).font = font_cell_bold
            ws_summary.cell(row=row_idx, column=2, value=name).font = font_cell
            ws_summary.cell(row=row_idx, column=3, value=score).font = font_cell_bold
            for c in range(1, 4):
                ws_summary.cell(row=row_idx, column=c).border = thin_border

        auto_fit_columns(ws_summary)

    # ------------------ SHEET: Staff Performance ------------------
    if report_type in ('360_master', 'staff_performance'):
        ws_staff = wb.create_sheet(title="Staff Performance")
        ws_staff["A1"] = f"{company.name} — Staff Outreach & Conversion Scorecard"
        ws_staff["A1"].font = font_title
        ws_staff["A2"] = f"Period: {period['start_date']} to {period['end_date']} | Active Employees: {len(data['staff_performance'])}"
        ws_staff["A2"].font = font_subtitle

        headers = [
            "Staff Member", "Designation", "Assigned Leads", "Total Follow-ups",
            "Calls Made", "WhatsApp", "Meetings", "Emails", "Deals Won", "Deals Lost",
            "Win Rate %", "Won Revenue (₹)", "Pipeline Value (₹)", "Overdue Follow-ups", "Last Activity"
        ]
        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws_staff.cell(row=row_idx, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_dark_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for s_idx, s in enumerate(data['staff_performance']):
            row_idx += 1
            fill = zebra_even_fill if s_idx % 2 == 0 else zebra_odd_fill
            row_data = [
                s['full_name'],
                s['designation'],
                s['total_assigned'],
                s['followups_count'],
                s['calls_count'],
                s['whatsapp_count'],
                s['meetings_count'],
                s['emails_count'],
                s['converted_count'],
                s['lost_count'],
                f"{s['conversion_rate']}%",
                float(s['won_revenue']),
                float(s['pipeline_value']),
                s['overdue_count'],
                s['last_activity_time'].strftime('%d %b %Y %H:%M') if s['last_activity_time'] else 'No Activity'
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws_staff.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_cell
                cell.fill = fill
                cell.border = thin_border
                if col_idx in (3, 4, 5, 6, 7, 8, 9, 10, 14):
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in (12, 13):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = '₹ #,##0.00'

        # Totals Row
        row_idx += 1
        ws_staff.cell(row=row_idx, column=1, value="TOTALS / SUMMARY").font = font_cell_bold
        ws_staff.cell(row=row_idx, column=1).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=1).border = thin_border
        ws_staff.cell(row=row_idx, column=2, value="").fill = total_row_fill

        ws_staff.cell(row=row_idx, column=3, value=sum(s['total_assigned'] for s in data['staff_performance'])).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=4, value=kpis['total_followups']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=5, value=kpis['total_calls']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=6, value=kpis['total_whatsapp']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=7, value=kpis['total_meetings']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=8, value=kpis['total_emails']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=9, value=kpis['total_converted']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=10, value=kpis['total_lost']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=11, value=f"{kpis['overall_conversion_rate']}%").fill = total_row_fill
        ws_staff.cell(row=row_idx, column=12, value=float(kpis['won_revenue'])).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=12).number_format = '₹ #,##0.00'
        ws_staff.cell(row=row_idx, column=13, value=float(kpis['active_pipeline_value'])).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=13).number_format = '₹ #,##0.00'
        ws_staff.cell(row=row_idx, column=14, value=kpis['overdue_count']).fill = total_row_fill
        ws_staff.cell(row=row_idx, column=15, value="").fill = total_row_fill

        for c in range(1, 16):
            ws_staff.cell(row=row_idx, column=c).border = thin_border
            ws_staff.cell(row=row_idx, column=c).font = font_cell_bold

        auto_fit_columns(ws_staff)

    # ------------------ SHEET: Lead Quality & Sources ------------------
    if report_type in ('360_master', 'lead_quality'):
        ws_sources = wb.create_sheet(title="Lead Quality & Sources")
        ws_sources["A1"] = f"{company.name} — Lead Acquisition Channels & Quality Matrix"
        ws_sources["A1"].font = font_title
        ws_sources["A2"] = f"Analyzes conversion rates, won revenue, and quality tiers per acquisition source"
        ws_sources["A2"].font = font_subtitle

        src_headers = [
            "Acquisition Source", "Total Leads", "Deals Won", "Deals Lost", "Open Leads",
            "Win Rate %", "Total Pipeline Value (₹)", "Won Revenue (₹)", "Avg Deal Size (₹)", "High Priority Leads", "Quality Tier"
        ]
        row_idx = 4
        for col_idx, h in enumerate(src_headers, 1):
            cell = ws_sources.cell(row=row_idx, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_dark_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for s_idx, src in enumerate(data['lead_sources']):
            row_idx += 1
            fill = zebra_even_fill if s_idx % 2 == 0 else zebra_odd_fill
            row_data = [
                src['source'],
                src['total_leads'],
                src['won_count'],
                src['lost_count'],
                src['open_count'],
                f"{src['conversion_rate']}%",
                float(src['total_value']),
                float(src['won_value']),
                float(src['avg_deal_size']),
                src['high_priority_count'],
                src['quality_tier']
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws_sources.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_cell
                cell.fill = fill
                cell.border = thin_border
                if col_idx in (2, 3, 4, 5, 6, 10):
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in (7, 8, 9):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = '₹ #,##0.00'

        # Priority Matrix Breakdown
        row_idx += 4
        ws_sources.cell(row=row_idx, column=1, value="LEAD PRIORITY DISTRIBUTION MATRIX").font = font_cell_bold
        row_idx += 1

        prio_headers = ["Priority Level", "Total Leads", "Won Deals", "Win Rate %", "Total Pipeline Value (₹)"]
        for col_idx, h in enumerate(prio_headers, 1):
            cell = ws_sources.cell(row=row_idx, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = sub_header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for p in data['priority_matrix']:
            row_idx += 1
            ws_sources.cell(row=row_idx, column=1, value=p['label']).font = font_cell_bold
            ws_sources.cell(row=row_idx, column=2, value=p['total_count']).alignment = Alignment(horizontal="center")
            ws_sources.cell(row=row_idx, column=3, value=p['won_count']).alignment = Alignment(horizontal="center")
            ws_sources.cell(row=row_idx, column=4, value=f"{p['conversion_rate']}%").alignment = Alignment(horizontal="center")
            c_val = ws_sources.cell(row=row_idx, column=5, value=float(p['total_value']))
            c_val.number_format = '₹ #,##0.00'
            c_val.alignment = Alignment(horizontal="right")
            for c in range(1, 6):
                ws_sources.cell(row=row_idx, column=c).border = thin_border

        auto_fit_columns(ws_sources)

    # ------------------ SHEET: Pipeline & Funnel ------------------
    if report_type in ('360_master', 'conversion_funnel'):
        ws_funnel = wb.create_sheet(title="Pipeline & Funnel")
        ws_funnel["A1"] = f"{company.name} — Sales Pipeline & Conversion Funnel"
        ws_funnel["A1"].font = font_title
        ws_funnel["A2"] = f"Track stage progression, pipeline values, and drop-offs across sales cycle"
        ws_funnel["A2"].font = font_subtitle

        funnel_headers = [
            "Sales Pipeline Stage",
            "Period Leads (Date-Filtered)",
            "Period Pipeline Value (₹)",
            "Cumulative Leads (All-Time Static)",
            "Cumulative Pipeline Value (₹)",
            "Share of Total %",
            "Avg Win Probability %"
        ]
        row_idx = 4
        for col_idx, h in enumerate(funnel_headers, 1):
            cell = ws_funnel.cell(row=row_idx, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_dark_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for f_idx, st in enumerate(data['pipeline_funnel']):
            row_idx += 1
            fill = zebra_even_fill if f_idx % 2 == 0 else zebra_odd_fill
            ws_funnel.cell(row=row_idx, column=1, value=st['status_name']).font = font_cell_bold
            ws_funnel.cell(row=row_idx, column=1).fill = fill
            
            c_p_cnt = ws_funnel.cell(row=row_idx, column=2, value=st['period_count'])
            c_p_cnt.alignment = Alignment(horizontal="center")
            c_p_cnt.fill = fill
            
            c_p_val = ws_funnel.cell(row=row_idx, column=3, value=float(st['period_value']))
            c_p_val.number_format = '₹ #,##0.00'
            c_p_val.alignment = Alignment(horizontal="right")
            c_p_val.fill = fill
            
            c_cnt = ws_funnel.cell(row=row_idx, column=4, value=st['count'])
            c_cnt.alignment = Alignment(horizontal="center")
            c_cnt.fill = fill
            
            c_val = ws_funnel.cell(row=row_idx, column=5, value=float(st['pipeline_value']))
            c_val.number_format = '₹ #,##0.00'
            c_val.alignment = Alignment(horizontal="right")
            c_val.fill = fill
            
            c_pct = ws_funnel.cell(row=row_idx, column=6, value=f"{st['percentage']}%")
            c_pct.alignment = Alignment(horizontal="center")
            c_pct.fill = fill
            
            c_prob = ws_funnel.cell(row=row_idx, column=7, value=f"{st['avg_probability']}%")
            c_prob.alignment = Alignment(horizontal="center")
            c_prob.fill = fill

            for c in range(1, 8):
                ws_funnel.cell(row=row_idx, column=c).border = thin_border

        # Lost Deals Reasons Breakdown
        row_idx += 4
        ws_funnel.cell(row=row_idx, column=1, value="LOST DEALS & DROP-OFF REASON ANALYSIS").font = font_cell_bold
        row_idx += 1

        lost_headers = ["Lost Reason / Friction Point", "Lost Deals Count", "Total Lost Value (₹)"]
        for col_idx, h in enumerate(lost_headers, 1):
            cell = ws_funnel.cell(row=row_idx, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = sub_header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        if data['lost_reasons']:
            for lr in data['lost_reasons']:
                row_idx += 1
                ws_funnel.cell(row=row_idx, column=1, value=lr['reason']).font = font_cell
                ws_funnel.cell(row=row_idx, column=2, value=lr['count']).alignment = Alignment(horizontal="center")
                c_val = ws_funnel.cell(row=row_idx, column=3, value=float(lr['lost_value']))
                c_val.number_format = '₹ #,##0.00'
                c_val.alignment = Alignment(horizontal="right")
                for c in range(1, 4):
                    ws_funnel.cell(row=row_idx, column=c).border = thin_border
        else:
            row_idx += 1
            ws_funnel.cell(row=row_idx, column=1, value="No lost deals recorded in this period.").font = font_cell
            ws_funnel.cell(row=row_idx, column=2, value=0)
            ws_funnel.cell(row=row_idx, column=3, value=0.00)
            for c in range(1, 4):
                ws_funnel.cell(row=row_idx, column=c).border = thin_border

        auto_fit_columns(ws_funnel)

    # ------------------ SHEET: Activity Audit Trail ------------------
    if report_type in ('360_master', 'activity_log', 'staff_performance'):
        ws_activity = wb.create_sheet(title="Activity Audit Trail")
        ws_activity["A1"] = f"{company.name} — Follow-ups & Activity Audit Trail"
        ws_activity["A1"].font = font_title
        ws_activity["A2"] = f"Chronological log of staff outreach interactions and customer touchpoints"
        ws_activity["A2"].font = font_subtitle

        act_headers = [
            "Date & Time", "Staff Member", "Lead Code", "Client / Company", "Contact Person",
            "Outreach Channel", "Interaction Summary", "Outcome", "Next Action", "Next Follow-up Due"
        ]
        row_idx = 4
        for col_idx, h in enumerate(act_headers, 1):
            cell = ws_activity.cell(row=row_idx, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_dark_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for a_idx, a in enumerate(data['activity_logs']):
            row_idx += 1
            fill = zebra_even_fill if a_idx % 2 == 0 else zebra_odd_fill
            row_data = [
                a['timestamp'].strftime('%d %b %Y %H:%M') if a['timestamp'] else '',
                a['staff_name'],
                a['lead_number'],
                a['company_name'],
                a['contact_person'],
                a['method'],
                a['summary'],
                a['outcome'],
                a['next_action'],
                a['next_followup_date'].strftime('%d %b %Y %H:%M') if a['next_followup_date'] else 'None'
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws_activity.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_cell
                cell.fill = fill
                cell.border = thin_border
                if col_idx in (1, 3, 6, 10):
                    cell.alignment = Alignment(horizontal="center")

        auto_fit_columns(ws_activity)

    # Remove the initial default empty sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_crm_pdf_report(company, report_type='360_master', start_date=None, end_date=None, staff_id=None, generated_by_user=None):
    """
    Generates an executive-ready, highly polished PDF report for CRM 360° Intelligence.
    Uses reportlab in landscape orientation with NumberedCanvas page numbering.
    """
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.pdfgen import canvas
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Spacer, Table, TableStyle, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    data = get_crm_360_analytics_data(company, start_date=start_date, end_date=end_date, staff_id=staff_id)
    kpis = data['summary_kpis']
    period = data['period_info']

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_decorations(num_pages)
                canvas.Canvas.showPage(self)
            canvas.Canvas.save(self)

        def draw_decorations(self, page_count):
            self.saveState()
            self.setFont("Helvetica", 7.5)
            self.setFillColor(colors.HexColor("#64748b"))

            # Running Top Header on pages > 1
            if self._pageNumber > 1:
                self.drawString(30, 595, f"{company.name} — CRM 360° Executive Intelligence Report")
                self.drawRightString(762, 595, f"Period: {period['start_date'].strftime('%d %b')} - {period['end_date'].strftime('%d %b %Y')}")
                self.setStrokeColor(colors.HexColor("#e2e8f0"))
                self.setLineWidth(0.5)
                self.line(30, 588, 762, 588)

            # Running Bottom Footer
            self.setStrokeColor(colors.HexColor("#e2e8f0"))
            self.setLineWidth(0.5)
            self.line(30, 25, 762, 25)
            gen_by_str = f" | Generated by: {generated_by_user.get_full_name() or generated_by_user.username}" if generated_by_user else ""
            self.drawString(30, 15, f"ByteBoot Business OS • CRM Performance Analytics • Confidential{gen_by_str}")
            self.drawRightString(762, 15, f"Page {self._pageNumber} of {page_count}")
            self.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=30,
        rightMargin=30,
        topMargin=32,
        bottomMargin=35
    )

    styles = getSampleStyleSheet()

    # Paragraph Styles
    title_p = ParagraphStyle('RTitle', fontName='Helvetica-Bold', fontSize=14, leading=17, textColor=colors.HexColor('#0f172a'))
    subtitle_p = ParagraphStyle('RSub', fontName='Helvetica', fontSize=8, leading=11, textColor=colors.HexColor('#64748b'))
    sec_header_p = ParagraphStyle('SecH', fontName='Helvetica-Bold', fontSize=10, leading=13, textColor=colors.HexColor('#0f172a'))
    
    th_p = ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=7.5, leading=9.5, textColor=colors.white, alignment=1)
    td_p = ParagraphStyle('TD', fontName='Helvetica', fontSize=7, leading=9, textColor=colors.HexColor('#1e293b'))
    td_bold_p = ParagraphStyle('TDBold', fontName='Helvetica-Bold', fontSize=7, leading=9, textColor=colors.HexColor('#0f172a'))
    td_center_p = ParagraphStyle('TDCenter', fontName='Helvetica', fontSize=7, leading=9, textColor=colors.HexColor('#1e293b'), alignment=1)
    td_right_p = ParagraphStyle('TDRight', fontName='Helvetica-Bold', fontSize=7, leading=9, textColor=colors.HexColor('#0f172a'), alignment=2)
    
    kpi_val_p = ParagraphStyle('KPIVal', fontName='Helvetica-Bold', fontSize=14, leading=16, alignment=1, textColor=colors.HexColor('#00a0da'))
    kpi_lbl_p = ParagraphStyle('KPILbl', fontName='Helvetica-Bold', fontSize=7, leading=8.5, alignment=1, textColor=colors.HexColor('#64748b'))
    kpi_sub_p = ParagraphStyle('KPISub', fontName='Helvetica', fontSize=6.5, leading=8, alignment=1, textColor=colors.HexColor('#94a3b8'))

    elements = []

    # 1. Header Banner Box
    scope_str = "All Company Staff" if not period['selected_staff'] else f"Staff: {period['selected_staff'].get_full_name() or period['selected_staff'].username}"
    report_title_map = {
        '360_master': '360° Comprehensive Sales & Staff Intelligence Audit',
        'staff_performance': 'Staff Outreach Velocity & Conversion Scorecard',
        'lead_quality': 'Lead Acquisition Channels & Quality Matrix',
        'conversion_funnel': 'Sales Pipeline Funnel & Stage Progression Audit',
        'activity_log': 'Follow-ups & Customer Outreach Audit Trail'
    }
    cur_title = report_title_map.get(report_type, 'CRM Analytics & Performance Intelligence Report')

    header_table_data = [
        [
            Paragraph(f"<b>{company.name.upper()}</b><br/><font size=12 color='#00a0da'><b>{cur_title}</b></font>", title_p),
            Paragraph(f"<b>Reporting Scope:</b> {scope_str}<br/><b>Date Window:</b> {period['start_date'].strftime('%d %b %Y')} to {period['end_date'].strftime('%d %b %Y')}<br/><b>Generated:</b> {period['generated_at'].strftime('%d %b %Y, %H:%M')}", subtitle_p)
        ]
    ]
    header_table = Table(header_table_data, colWidths=[480, 252])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LINEBELOW', (0,0), (-1,-1), 1.5, colors.HexColor('#00a0da')),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 8))

    # 2. Executive KPI Metrics Cards
    kpi_data = [
        [
            Paragraph(str(kpis['total_followups']), kpi_val_p),
            Paragraph(str(kpis['total_calls']), ParagraphStyle('k2', parent=kpi_val_p, textColor=colors.HexColor('#10b981'))),
            Paragraph(str(kpis['total_converted']), ParagraphStyle('k3', parent=kpi_val_p, textColor=colors.HexColor('#f59e0b'))),
            Paragraph(f"{kpis['overall_conversion_rate']}%", ParagraphStyle('k4', parent=kpi_val_p, textColor=colors.HexColor('#8b5cf6'))),
            Paragraph(f"₹ {float(kpis['won_revenue']):,.0f}", ParagraphStyle('k5', parent=kpi_val_p, textColor=colors.HexColor('#059669'))),
            Paragraph(f"₹ {float(kpis['active_pipeline_value']):,.0f}", ParagraphStyle('k6', parent=kpi_val_p, textColor=colors.HexColor('#00a0da'))),
        ],
        [
            Paragraph("TOTAL FOLLOW-UPS", kpi_lbl_p),
            Paragraph("DIRECT CALLS", kpi_lbl_p),
            Paragraph("DEALS CONVERTED", kpi_lbl_p),
            Paragraph("WIN RATE %", kpi_lbl_p),
            Paragraph("WON REVENUE", kpi_lbl_p),
            Paragraph("ACTIVE PIPELINE", kpi_lbl_p),
        ],
        [
            Paragraph(f"{kpis['total_staff']} active staff", kpi_sub_p),
            Paragraph(f"{kpis['total_whatsapp']} WhatsApp msgs", kpi_sub_p),
            Paragraph(f"{kpis['total_lost']} deals lost", kpi_sub_p),
            Paragraph(f"Win/Closed: {kpis['pipeline_win_rate']}%", kpi_sub_p),
            Paragraph(f"Avg Deal: ₹ {float(kpis['avg_deal_size']):,.0f}", kpi_sub_p),
            Paragraph(f"{kpis['overdue_count']} overdue pending", kpi_sub_p),
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[122, 122, 122, 122, 122, 122])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 10))

    # 3. Staff Performance Table
    if report_type in ('360_master', 'staff_performance'):
        elements.append(Paragraph("1. Staff Outreach Velocity & Conversion Scorecard", sec_header_p))
        elements.append(Spacer(1, 4))

        staff_rows = [
            [
                Paragraph("Staff Member", th_p),
                Paragraph("Role", th_p),
                Paragraph("Assigned", th_p),
                Paragraph("Follow-ups", th_p),
                Paragraph("Calls", th_p),
                Paragraph("WhatsApp/Other", th_p),
                Paragraph("Won", th_p),
                Paragraph("Lost", th_p),
                Paragraph("Win Rate", th_p),
                Paragraph("Overdue", th_p),
                Paragraph("Won Revenue (₹)", th_p),
                Paragraph("Active Pipeline (₹)", th_p),
            ]
        ]
        for s in data['staff_performance']:
            other_ch = s['whatsapp_count'] + s['meetings_count'] + s['emails_count']
            staff_rows.append([
                Paragraph(f"<b>{s['full_name']}</b><br/><font color='#64748b'>@{s['username']}</font>", td_p),
                Paragraph(s['designation'][:16], td_center_p),
                Paragraph(str(s['total_assigned']), td_center_p),
                Paragraph(f"<b>{s['followups_count']}</b>", td_center_p),
                Paragraph(str(s['calls_count']), td_center_p),
                Paragraph(str(other_ch), td_center_p),
                Paragraph(f"<b>{s['converted_count']}</b>", td_center_p),
                Paragraph(str(s['lost_count']), td_center_p),
                Paragraph(f"<b>{s['conversion_rate']}%</b>", td_center_p),
                Paragraph(f"<font color='{'#dc2626' if s['overdue_count'] > 0 else '#059669'}'><b>{s['overdue_count']}</b></font>", td_center_p),
                Paragraph(f"₹ {float(s['won_revenue']):,.0f}", td_right_p),
                Paragraph(f"₹ {float(s['pipeline_value']):,.0f}", td_right_p),
            ])

        staff_col_w = [115, 65, 45, 50, 40, 55, 35, 35, 48, 45, 99, 100]
        st_table = Table(staff_rows, colWidths=staff_col_w, repeatRows=1)
        st_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f172a')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')])
        ]))
        elements.append(st_table)

    # 4. Lead Quality Channels & Pipeline Funnel
    if report_type in ('360_master', 'lead_quality', 'conversion_funnel'):
        if report_type == '360_master':
            elements.append(PageBreak())

        elements.append(Paragraph("2. Lead Acquisition Channels & Quality Matrix", sec_header_p))
        elements.append(Spacer(1, 4))

        src_rows = [
            [
                Paragraph("Source Channel", th_p),
                Paragraph("Total Leads", th_p),
                Paragraph("Deals Won", th_p),
                Paragraph("Deals Lost", th_p),
                Paragraph("Open Leads", th_p),
                Paragraph("Win Rate %", th_p),
                Paragraph("Total Value (₹)", th_p),
                Paragraph("Won Revenue (₹)", th_p),
                Paragraph("High Priority", th_p),
                Paragraph("Quality Rating", th_p),
            ]
        ]
        for src in data['lead_sources']:
            src_rows.append([
                Paragraph(f"<b>{src['source']}</b>", td_p),
                Paragraph(str(src['total_leads']), td_center_p),
                Paragraph(str(src['won_count']), td_center_p),
                Paragraph(str(src['lost_count']), td_center_p),
                Paragraph(str(src['open_count']), td_center_p),
                Paragraph(f"<b>{src['conversion_rate']}%</b>", td_center_p),
                Paragraph(f"₹ {float(src['total_value']):,.0f}", td_right_p),
                Paragraph(f"₹ {float(src['won_value']):,.0f}", td_right_p),
                Paragraph(str(src['high_priority_count']), td_center_p),
                Paragraph(src['quality_tier'], td_center_p),
            ])

        src_col_w = [110, 50, 45, 45, 45, 55, 95, 95, 60, 132]
        src_table = Table(src_rows, colWidths=src_col_w, repeatRows=1)
        src_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f172a')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')])
        ]))
        elements.append(src_table)
        elements.append(Spacer(1, 10))

        # Pipeline Funnel Progression
        elements.append(Paragraph("3. Sales Pipeline Funnel & Stage Progression", sec_header_p))
        elements.append(Spacer(1, 4))

        fnl_rows = [
            [
                Paragraph("Pipeline Stage", th_p),
                Paragraph("Period Leads", th_p),
                Paragraph("Period Value (₹)", th_p),
                Paragraph("Cumulative Leads", th_p),
                Paragraph("Cumulative Value (₹)", th_p),
                Paragraph("Share %", th_p),
                Paragraph("Win Prob %", th_p),
            ]
        ]
        for st in data['pipeline_funnel']:
            fnl_rows.append([
                Paragraph(f"<b>{st['status_name']}</b>", td_p),
                Paragraph(str(st['period_count']), td_center_p),
                Paragraph(f"₹ {float(st['period_value']):,.0f}", td_right_p),
                Paragraph(str(st['count']), td_center_p),
                Paragraph(f"₹ {float(st['pipeline_value']):,.0f}", td_right_p),
                Paragraph(f"{st['percentage']}%", td_center_p),
                Paragraph(f"{st['avg_probability']}%", td_center_p),
            ])

        fnl_col_w = [145, 80, 120, 95, 130, 62, 100]
        fnl_table = Table(fnl_rows, colWidths=fnl_col_w, repeatRows=1)
        fnl_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')])
        ]))
        elements.append(fnl_table)
        elements.append(Spacer(1, 10))

        # Executive Insights Box
        top_src = max(data['lead_sources'], key=lambda s: s['won_value']) if data['lead_sources'] and any(s['won_value'] > 0 for s in data['lead_sources']) else None
        top_rate_src = max(data['lead_sources'], key=lambda s: s['conversion_rate']) if data['lead_sources'] and any(s['conversion_rate'] > 0 for s in data['lead_sources']) else None
        insights_text = f"""
        <b>Executive Intelligence Highlights:</b><br/>
        • <b>Top Converting Channel:</b> {top_rate_src['source'] if top_rate_src else 'N/A'} with a win rate of <b>{top_rate_src['conversion_rate'] if top_rate_src else '0'}%</b>.<br/>
        • <b>Highest Revenue Generator:</b> {top_src['source'] if top_src else 'N/A'} yielding <b>₹ {float(top_src['won_value']):,.2f}</b> in converted customer value.<br/>
        • <b>Follow-up Compliance:</b> {kpis['overdue_count']} leads currently overdue across the sales team. Prioritize resolving backlog to prevent deal attrition.
        """
        insight_table = Table([[Paragraph(insights_text, td_p)]], colWidths=[732])
        insight_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#eff6ff')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#93c5fd')),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
            ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ]))
        elements.append(insight_table)

    # 5. Activity Log (if specific report or detailed requested)
    if report_type == 'activity_log':
        elements.append(Paragraph("Customer Follow-ups & Activity Audit Trail", sec_header_p))
        elements.append(Spacer(1, 4))

        act_rows = [
            [
                Paragraph("Date & Time", th_p),
                Paragraph("Staff", th_p),
                Paragraph("Lead Code & Client", th_p),
                Paragraph("Channel", th_p),
                Paragraph("Interaction Summary & Outcome", th_p),
                Paragraph("Next Action", th_p),
            ]
        ]
        for a in data['activity_logs']:
            act_rows.append([
                Paragraph(a['timestamp'].strftime('%d %b %H:%M') if a['timestamp'] else '', td_center_p),
                Paragraph(a['staff_name'], td_p),
                Paragraph(f"<b>{a['lead_number']}</b><br/>{a['company_name']}", td_p),
                Paragraph(a['method'], td_center_p),
                Paragraph(f"{a['summary']}<br/><font color='#059669'><b>Outcome:</b> {a['outcome']}</font>", td_p),
                Paragraph(a['next_action'], td_p),
            ])

        act_col_w = [75, 85, 115, 65, 272, 120]
        act_table = Table(act_rows, colWidths=act_col_w, repeatRows=1)
        act_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f172a')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')])
        ]))
        elements.append(act_table)

    doc.build(elements, canvasmaker=NumberedCanvas)
    buf.seek(0)
    return buf

